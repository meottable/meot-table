import datetime as dt
import importlib.util
import tempfile
import unittest
from pathlib import Path
import openpyxl

spec = importlib.util.spec_from_file_location('builder', Path(__file__).parents[1] / 'ledger/tools/build_source_snapshot.py')
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


class ProjectionTests(unittest.TestCase):
    def book(self):
        book = openpyxl.Workbook()
        book.remove(book.active)
        order = {'주문일': dt.date(2026,10,15), '견적번호':'QA-1', '고객/상호':'검증용', '규격':'테스트', '등급':'테스트', '테이블':1, '의자':0, '수저통':0, '총 계약금액':1000, '계약금 입금':None, '잔금 입금':None, '총입금':None, '미수금':None, '결제상태':'입금 미확인', '생산상태':'주문', '납품상태':'미정', '비고':'검증용', '정산 차감액':None}
        sale = {'거래일자':dt.date(2026,10,15), '견적번호':'QA-1', '합계금액':1000, '입금액':None, '미수금':None, '결제상태':'입금 미확인'}
        for name, row in [('주문관리',order), ('매출입력',sale)]:
            s = book.create_sheet(name); s.append(list(row)); s.append(list(row.values()))
        book.create_sheet('모바일입력').append(['날짜','입출금','금액','분류','내용/거래처','메모'])
        book.create_sheet('영업관리').append(['월 목표'])
        return book

    def set(self, book, sheet, field, value):
        s=book[sheet]; headers=[c.value for c in s[1]]; s.cell(2,headers.index(field)+1,value)

    def project(self, book):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'synthetic.xlsx'; book.save(p); return builder.projection(p)

    def test_unknown_is_not_zero_or_a_receipt(self):
        p=self.project(self.book()); o=p['orders'][0]
        self.assertEqual(p['schema'],3); self.assertEqual(o['amount'],1000); self.assertEqual(o['paymentStatus'],'unknown')
        self.assertIsNone(o['paid']); self.assertIsNone(o['balance']); self.assertEqual(p['cash'],[])

    def test_unknown_with_any_payment_number_is_rejected(self):
        for sheet,field in [('주문관리','계약금 입금'),('주문관리','총입금'),('매출입력','입금액')]:
            with self.subTest(sheet=sheet,field=field):
                b=self.book(); self.set(b,sheet,field,0)
                with self.assertRaises(ValueError): self.project(b)

    def test_missing_payment_cache_without_unknown_status_is_rejected(self):
        b=self.book(); self.set(b,'주문관리','결제상태','미입금')
        with self.assertRaises(ValueError): self.project(b)

    def test_explicit_zero_and_balance_can_publish(self):
        b=self.book()
        for sheet,paid in [('주문관리','총입금'),('매출입력','입금액')]:
            self.set(b,sheet,'결제상태','미입금'); self.set(b,sheet,paid,0); self.set(b,sheet,'미수금',1000)
        o=self.project(b)['orders'][0]; self.assertEqual((o['paymentStatus'],o['paid'],o['balance']),('known',0,1000))

    def test_different_status_between_sheets_is_rejected(self):
        b=self.book(); self.set(b,'매출입력','결제상태','미입금'); self.set(b,'매출입력','입금액',0); self.set(b,'매출입력','미수금',1000)
        with self.assertRaises(ValueError): self.project(b)


if __name__=='__main__': unittest.main()
